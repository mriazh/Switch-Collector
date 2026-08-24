#!/usr/bin/env python3
"""
Unified Switch Collector Tool
Merges and unifies collection of Serial Numbers (SN), PoE Remaining Power, and Huawei LLDP Remote Device counts.
"""

import os
import re
import sys
import argparse
import logging
import time
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from netmiko import ConnectHandler
from netmiko.exceptions import NetmikoTimeoutException, NetmikoAuthenticationException, ReadTimeout

# Where the operator's switch.txt, outputs\ and logs\ live.
#
# In a frozen PyInstaller build, __file__ points inside the _internal runtime
# directory, so using it directly would make the executable look for switch.txt
# in _internal\ instead of the folder the operator actually extracted. When
# frozen, the real base directory is the folder holding the .exe.
if getattr(sys, "frozen", False):
    BASE_DIR = Path(sys.executable).resolve().parent
else:
    BASE_DIR = Path(__file__).resolve().parent

# Operator-facing remediation hints surfaced in the Solution Hint column.
HINT_AUTH = ("Verify username & password in switch.txt, ensure exec privilege is "
             "enabled, or check TACACS/RADIUS server.")
HINT_TIMEOUT = ("Check ping connectivity to switch IP, verify device is powered on, "
                "and ensure port 22 is allowed by firewall/ACL.")
HINT_CONNECTION_REFUSED = ("SSH service is inactive on switch. Connect via console/telnet "
                           "and run ip ssh.")
HINT_PARSE = ("Command output did not match expected ArubaOS format. Check firmware "
              "version or inspect raw output in logs.")
HINT_FAN = "Cooling fan failure on physical member. Inspect chassis fan modules in rack."
HINT_PARTIAL = ("Some collector modules failed to return data. Check hardware feature "
                "availability or increase read-timeout.")

def parse_args():
    """
    Parses command-line arguments.
    """
    parser = argparse.ArgumentParser(description="Unified Switch Collector Tool")
    parser.add_argument(
        "--mode",
        required=False,
        default=None,
        choices=["sn", "fan", "poe", "huawei-count", "all"],
        help="Collection mode: sn, fan, poe, huawei-count, all. "
             "Omit to use the interactive menu."
    )
    parser.add_argument(
        "--inventory",
        default=None,
        help="Path to switch inventory file (default: switch.txt in script dir)"
    )
    parser.add_argument(
        "--output-dir",
        default=None,
        help="Path to output directory (default: outputs/ in script dir)"
    )
    parser.add_argument(
        "--device-types",
        default="hp_procurve,aruba_os",
        help="Comma-separated Netmiko device types to try (default: hp_procurve,aruba_os)"
    )
    parser.add_argument(
        "--read-timeout",
        type=int,
        default=45,
        help="Command read timeout in seconds (default: 45)"
    )
    parser.add_argument(
        "--workers",
        type=int,
        default=None,
        help="Number of parallel workers (default: 5, or prompt in interactive mode). "
             "Use 1 for sequential execution."
    )
    parser.add_argument(
        "--log-dir",
        default="logs",
        help="Directory for log files (default: logs)"
    )
    return parser.parse_args()


# ==================== INTERACTIVE MENU ====================

# Menu option -> collection mode.
MENU_OPTIONS = {
    "1": "all",
    "2": "fan",
    "3": "sn",
    "4": "poe",
    "5": "huawei-count",
}

MENU_BANNER = """===================================================
               SWITCH COLLECTOR MASTER
===================================================
 [1] Full Master Audit (SN, Fan, PoE, Huawei AP)
 [2] Fan Health Audit Only
 [3] Hardware Asset Audit (Serial Number & ROM)
 [4] PoE Capacity Audit Only
 [5] Huawei AP Migration Audit Only
 [0] Exit
==================================================="""

DEFAULT_WORKERS = 5


def print_menu():
    """Prints the interactive selection banner."""
    print(MENU_BANNER)


def prompt_mode(input_fn=None, print_fn=None):
    """
    Presents the menu and returns the selected collection mode.

    Returns None when the user picks Exit. Re-prompts on invalid input so a
    typo does not abort an audit.
    """
    # Resolved at call time so tests can patch builtins.input / print.
    input_fn = input_fn or input
    print_fn = print_fn or print

    while True:
        print_fn(MENU_BANNER)
        choice = input_fn("Select an option [0-5]: ").strip()

        if choice == "0":
            return None
        if choice in MENU_OPTIONS:
            return MENU_OPTIONS[choice]

        print_fn("Invalid selection. Please enter a number between 0 and 5.")


def prompt_workers(default=DEFAULT_WORKERS, input_fn=None, print_fn=None):
    """
    Prompts for the worker thread count.

    An empty response keeps the default. A non-numeric value is rejected so the
    caller never receives a nonsensical thread count.
    """
    input_fn = input_fn or input
    print_fn = print_fn or print

    while True:
        raw = input_fn("Enter number of worker threads [default {}]: ".format(default)).strip()
        if not raw:
            return default
        try:
            value = int(raw)
        except ValueError:
            print_fn("Invalid input. Please enter a whole number.")
            continue
        if value < 1:
            print_fn("Worker count must be at least 1.")
            continue
        return value


def resolve_runtime_options(args, input_fn=None, print_fn=None):
    """
    Fills in mode and workers from the interactive menu when they were omitted.

    Headless usage (any --mode supplied) bypasses the menu entirely.
    Returns (mode, workers); mode is None when the user chose Exit.
    """
    workers = getattr(args, "workers", None)

    if getattr(args, "mode", None):
        # Headless: no prompts, fall back to the default thread count.
        return args.mode, (workers if workers else DEFAULT_WORKERS)

    mode = prompt_mode(input_fn=input_fn, print_fn=print_fn)
    if mode is None:
        return None, None

    # An explicit --workers wins; otherwise ask interactively.
    if workers is None:
        workers = prompt_workers(default=DEFAULT_WORKERS, input_fn=input_fn, print_fn=print_fn)
    return mode, workers

def parse_inventory(path):
    """
    Parses the inventory file.
    Each non-empty line must have exactly 4 fields separated by '|'.
    Rejects any empty fields (label, ip, username, or password).
    Returns (records, errors) where records is a list of valid dictionaries
    and errors is a list of error dictionaries for any malformed/invalid lines.
    """
    records = []
    errors = []
    if not os.path.exists(path):
        raise FileNotFoundError(f"Inventory file not found: {path}")

    with open(path, "r", encoding="utf-8") as f:
        for line_num, line in enumerate(f, 1):
            stripped = line.strip()
            if not stripped:
                continue
            
            # Clean non-breaking spaces and split by pipe
            parts = [p.replace('\xa0', ' ').strip() for p in line.split('|')]
            if len(parts) != 4:
                # Malformed entry (incorrect field count)
                label = parts[0] if len(parts) > 0 and parts[0] else f"Line {line_num}"
                ip = parts[1] if len(parts) > 1 and parts[1] else "Unknown"
                errors.append({
                    "label": label,
                    "ip": ip,
                    "error_type": "InvalidInput",
                    "error_message": f"Line {line_num} does not have exactly 4 fields (got {len(parts)})",
                    "attempts": 0
                })
            else:
                label, ip, username, password = parts
                # Reject if any field is empty
                if not label or not ip or not username or not password:
                    missing_fields = []
                    if not label: missing_fields.append("label")
                    if not ip: missing_fields.append("ip")
                    if not username: missing_fields.append("username")
                    if not password: missing_fields.append("password")
                    
                    label_val = label if label else f"Line {line_num}"
                    ip_val = ip if ip else "Unknown"
                    errors.append({
                        "label": label_val,
                        "ip": ip_val,
                        "error_type": "InvalidInput",
                        "error_message": f"Line {line_num} has empty required fields: {', '.join(missing_fields)}",
                        "attempts": 0
                    })
                else:
                    records.append({
                        "label": label,
                        "ip": ip,
                        "username": username,
                        "password": password
                    })
    return records, errors

def _extract_field(text, patterns):
    """
    Returns the first value matching any of `patterns`, or "".
    Used to read the same logical field across ProCurve and ArubaOS-CX syntax.
    """
    for pattern in patterns:
        match = re.search(pattern, text, re.IGNORECASE)
        if match:
            value = (match.group(1) or "").strip()
            if value:
                return value
    return ""


# ProCurve (ArubaOS-S) and ArubaOS-CX label variants for the same field.
HOSTNAME_PATTERNS = [r'System Name[ \t]*:[ \t]*(\S+)', r'Hostname[ \t]*:[ \t]*(\S+)']
SERIAL_PATTERNS = [r'Chassis Serial Nbr[ \t]*:[ \t]*(\S+)', r'Serial Number[ \t]*:[ \t]*(\S+)']
ROM_PATTERNS = [
    r'ROM Version[ \t]*:[ \t]*(\S+)',
    r'ArubaOS-CX Version[ \t]*:[ \t]*(\S+)',
    r'System Description[ \t]*:[ \t]*(\S+)',
]


def parse_system_information(output, label=""):
    """
    Parses system information from ProCurve (`show system information`) and
    ArubaOS-CX (`show system`).

    Hostname comes from `System Name` or `Hostname`, Serial Number from
    `Serial Number` or `Chassis Serial Nbr`, and the firmware version from
    `ROM Version`, `ArubaOS-CX Version` or `System Description`.

    If VSF-Member blocks exist, extracts ROM Version and Serial Number for each.
    If no VSF-Member exists, extracts one global ROM Version and Serial Number.
    Raises ValueError if any ROM Version or Serial Number is missing (only for standalone).
    Returns a list of dicts: [{'hostname': ..., 'serial_number': ..., 'rom_version': ...}]
    """
    # 1. Extract hostname (System Name / Hostname), with fallback to inventory label
    hostname = _extract_field(output, HOSTNAME_PATTERNS) or label

    # 2. Check for VSF-Member blocks
    vsf_matches = list(re.finditer(r'VSF-Member\s*:\s*(\d+)', output, re.IGNORECASE))

    results = []
    if vsf_matches:
        for i, match in enumerate(vsf_matches):
            member_id = match.group(1)
            start_idx = match.end()
            end_idx = vsf_matches[i+1].start() if i + 1 < len(vsf_matches) else len(output)
            member_block = output[start_idx:end_idx]

            rom_version = _extract_field(member_block, ROM_PATTERNS)
            serial_number = _extract_field(member_block, SERIAL_PATTERNS)

            # VSF members are allowed to have blank Serial/ROM without throwing an error
            results.append({
                "hostname": hostname,
                "serial_number": serial_number,
                "rom_version": rom_version
            })
    else:
        # Standalone switch (ProCurve or ArubaOS-CX chassis)
        rom_version = _extract_field(output, ROM_PATTERNS)
        serial_number = _extract_field(output, SERIAL_PATTERNS)

        if not rom_version or not serial_number:
            missing = []
            if not rom_version: missing.append("ROM Version")
            if not serial_number: missing.append("Serial Number")
            raise ValueError(
                f"Missing {', '.join(missing)} in standalone system information"
            )

        results.append({
            "hostname": hostname,
            "serial_number": serial_number,
            "rom_version": rom_version
        })

    return results

NON_POE = "Non-PoE"

# Markers that mean the platform rejected the PoE command outright.
NON_POE_MARKERS = (
    "invalid input", "unknown command", "invalid command", "not supported",
    "no poe", "non-poe", "does not support", "not available on this",
    "% Invalid", "incomplete command",
)


def is_non_poe_output(output):
    """
    True when the CLI rejected `show power-over-ethernet`.

    Core chassis such as the Aruba 8400X have no PoE subsystem, so an
    unsupported command is a normal result, not a collection failure.
    """
    if not output:
        return False
    lowered = output.lower()
    return any(marker in lowered for marker in NON_POE_MARKERS)


def parse_power_output(output):
    """
    Parses 'show power-over-ethernet' output.
    Looks for 'Total Remaining Power' and returns 'xx W'.
    Returns 'Non-PoE' when the platform has no PoE support.
    If missing, returns empty string.
    """
    if is_non_poe_output(output):
        return NON_POE

    for line in output.splitlines():
        line = line.strip()
        if "Total Remaining Power" in line:
            # Pattern: digits optionally followed by spaces and W
            match = re.search(r'(\d+)[ \t]*W', line, re.IGNORECASE)
            if match:
                return f"{match.group(1)} W"
            # Fallback: take what's after the colon
            if ':' in line:
                parts = line.split(':', 1)
                if len(parts) == 2:
                    val = parts[1].strip()
                    if val:
                        return val
    return ""

def get_collectors_to_run(mode):
    """
    Returns the ordered list of collector names for a given --mode.
    'all' runs every collector, including 'fan'.
    """
    if mode == "all":
        return ["sn", "fan", "poe", "huawei-count"]
    return [mode]


COLLECTOR_COMMANDS = {
    "sn": "show system information",
    "fan": "show system fans",
    "poe": "show power-over-ethernet",
    "huawei-count": "show lldp info remote-device",
}

# ArubaOS-CX equivalents tried when the ProCurve command is rejected.
COLLECTOR_FALLBACK_COMMANDS = {
    "sn": "show system",
    "fan": "show environment fan",
    "huawei-count": "show lldp neighbor-info",
}

# Markers that mean the CLI rejected a command's syntax.
INVALID_COMMAND_MARKERS = (
    "invalid input", "invalid command", "unknown command", "incomplete command",
    "% invalid",
)


def parse_fan_output(output):
    """
    Parses 'show system fans' output.

    VSF stacks: each member reports 'X / Y Fans in Failure State'. Operational
    fans are (total - failures)/total, e.g. '0 / 4' -> '4/4', '1 / 4' -> '3/4'.
    Standalone switches: same ratio from a single global line or the fan table.

    Returns a dict keyed by VSF member id as a string:
      VSF stack     -> {'2': '4/4', '3': '4/4', '4': '4/4'}
      Standalone    -> {'1': '2/2'}
      Fanless       -> {'1': 'Fanless'}
      Nothing found -> {'1': 'N/A'}
    """
    if not output or not output.strip():
        return {"1": "N/A"}

    # Fanless switches report that they carry no fans at all.
    if re.search(
        r'fanless|no fans\b|does not contain any fans?\b|not equipped with fans?\b|'
        r'no cooling fans?\b',
        output, re.IGNORECASE
    ):
        return {"1": "Fanless"}

    result = {}

    # ArubaOS-CX: 'show environment fan' lists M/slot/fan rows with a status.
    cx = _parse_cx_fan_table(output)
    if cx:
        return cx

    # VSF stacks carry a 'VSF-Member :<id>' header per member block.
    vsf_matches = list(re.finditer(r'VSF-Member\s*:?\s*(\d+)', output, re.IGNORECASE))
    if vsf_matches:
        for i, match in enumerate(vsf_matches):
            member_id = match.group(1)
            start_idx = match.end()
            end_idx = vsf_matches[i + 1].start() if i + 1 < len(vsf_matches) else len(output)
            member_block = output[start_idx:end_idx]
            status = _extract_fan_ratio(member_block)
            if status:
                result[member_id] = status
        if result:
            return result

    # Standalone: use a single global ratio.
    status = _extract_fan_ratio(output)
    if status:
        return {"1": status}

    return {"1": "N/A"}


# ArubaOS-CX fan statuses that count as a healthy, working fan.
CX_FAN_HEALTHY = ("ok", "ready")
CX_FAN_FAILED = ("failed", "fault", "down")

# '1/1/1  <sensor>  ok  <rpm>' style row from 'show environment fan'.
CX_FAN_ROW = re.compile(
    r'^\s*(\d+)\s*/\s*(\d+)\s*/\s*(\d+)\s+'          # member / slot / fan
    r'.*?\b(ok|ready|failed|fault|down)\b\s*'         # status word
    r'\S*\d*.*$',                                     # trailing reading / duty
    re.IGNORECASE,
)


def _parse_cx_fan_table(text):
    """
    Parses the ArubaOS-CX `show environment fan` table.

    Each row looks like '1/1/1  <sensor>  ok  <rpm>'. Fans reporting 'ok' or
    'ready' count as working. Results are bucketed per VSF member using the
    first digit of the M/slot/fan path, e.g. 18 ok fans -> {'1': '18/18'}.

    Returns {} when the table is not present.
    """
    if not text or not re.search(r'fan\s+information', text, re.IGNORECASE):
        return {}

    working = {}
    total = {}
    for line in text.splitlines():
        match = CX_FAN_ROW.match(line)
        if not match:
            continue
        member_id = match.group(1)
        status = match.group(4).lower()
        total[member_id] = total.get(member_id, 0) + 1
        if status in CX_FAN_HEALTHY:
            working[member_id] = working.get(member_id, 0) + 1

    return {
        member: "{}/{}".format(working.get(member, 0), count)
        for member, count in total.items()
        if count > 0
    }


def _extract_fan_ratio(text):
    """
    Extracts an operational/total fan ratio from a block of fan output.
    Returns e.g. '4/4' or None when no fan information is present.
    """
    # '(\d+) / (\d+) Fans in Failure State' => failures / total
    match = re.search(
        r'(\d+)\s*[/\\]\s*(\d+)\s*Fans?\s+in\s+Failure\s+State',
        text, re.IGNORECASE
    )
    if match:
        failures = int(match.group(1))
        total = int(match.group(2))
        if total <= 0:
            return None
        return f"{total - failures}/{total}"

    # Individual fan table rows, e.g. 'Fan 1  OK' / 'Fan 2  Failed'
    rows = re.findall(r'^\s*Fan\s*(\d+)\s*[:|-]?\s*(OK|Failed|Not Present)?', text,
                      re.IGNORECASE | re.MULTILINE)
    present = [(num, state) for num, state in rows if state and state.lower() != "not present"]
    if present:
        total = len(present)
        failures = sum(1 for _, state in present if state.lower() == "failed")
        return f"{total - failures}/{total}"

    return None


# A row that begins with a port path: '1/1/3', '1/5' or a bare '1'.
PORT_PREFIX = re.compile(r'^\d+\s*/|^\d+\s*\|')


def count_huawei_devices(output):
    """
    Counts LLDP neighbours containing 'HUAWEI' (case-insensitive), attributed
    per VSF physical member.

    Handles both table layouts:
      ProCurve `show lldp info remote-device`  -> pipe-delimited rows after a
        dashed header separator; LocalPort is 'M/P' (e.g. '1/5').
      ArubaOS-CX `show lldp neighbor-info`     -> space-delimited rows that
        begin with a port path 'M/Slot/Port' (e.g. '1/1/3'), with no separator
        line required.

    A port without a leading integer ('1', 'A1') belongs to member 1.

    Returns a dict of member id -> count, e.g. {'1': 2, '2': 1}.
    """
    counts = {}
    in_data = False

    for line in output.splitlines():
        stripped = line.strip()
        # Header separator starts the data region in the ProCurve layout.
        if stripped.startswith('---') or stripped.startswith('+'):
            in_data = True
            continue
        if not stripped:
            continue

        # The ArubaOS-CX layout has no separator, so a leading port path is
        # enough to identify a data row.
        if not in_data and not PORT_PREFIX.match(stripped):
            continue

        if not re.search(r'HUAWEI', stripped, re.IGNORECASE):
            continue

        member_id = _member_from_local_port(stripped)
        counts[member_id] = counts.get(member_id, 0) + 1

    return counts


def _member_from_local_port(line):
    """
    Derives the VSF member id from an LLDP neighbour table row.

    Handles both table layouts:
      ProCurve 'show lldp info remote-device' -> pipe-separated, first column
        is the LocalPort ('1/5', 'A1').
      ArubaOS-CX 'show lldp neighbor-info'    -> space-separated, row starts
        with Member/Slot/Port ('1/1/3', '1/2/4').

    'M/...' and 'M/S/P' both resolve to member M; a value with no leading
    integer ('1', 'A1') resolves to member 1.
    """
    # Prefer the first pipe-delimited column, else the leading token.
    if '|' in line:
        local_port = line.split('|')[0].strip()
    else:
        local_port = line.split()[0].strip() if line.split() else ""

    # '1/5' or '1/1/3' -> member '1'
    match = re.match(r'^(\d+)\s*/', local_port)
    if match:
        return match.group(1)
    return "1"


def get_solution_hint(error_type, error_msg, fan_status=None):
    """
    Returns an operator-facing remediation hint for a given failure.

    Hints are matched in order of specificity: authentication, connectivity,
    parse failures, then fan health, then generic partial-failure guidance.
    """
    error_type = (error_type or "").lower()

    if "authentication" in error_type or "NetmikoAuthenticationException" in error_type:
        return HINT_AUTH

    if "timeout" in error_type or "NetmikoTimeoutException" in error_type or "ReadTimeout" in error_type:
        return HINT_TIMEOUT

    if "ConnectionRefused" in error_type or "refused" in error_type:
        return HINT_CONNECTION_REFUSED

    if "ParseError" in error_type or "parse" in error_type:
        return HINT_PARSE

    # Fan degradation: '3/4' means one fan is down.
    if fan_status and fan_is_degraded(fan_status):
        return HINT_FAN

    if "partial" in error_type or "warning" in error_type:
        return HINT_PARTIAL

    return ""


def _fan_status_values(fan_status):
    """Normalises a fan status (dict or scalar) into an iterable of values."""
    if isinstance(fan_status, dict):
        return list(fan_status.values())
    if isinstance(fan_status, (list, tuple, set)):
        return list(fan_status)
    return [fan_status]


def fan_ratio(value):
    """
    Parses a fan status such as '3/4' into (working, total).
    Returns None for 'Fanless', 'N/A' or anything unparseable.
    """
    match = re.match(r'^\s*(\d+)\s*/\s*(\d+)\s*$', str(value or ""))
    if not match:
        return None
    return int(match.group(1)), int(match.group(2))


def fan_is_degraded(value):
    """
    True when a fan status shows fewer working fans than total fans.

    '4/4' -> False (healthy); '3/4' and '1/2' -> True (degraded).
    'Fanless' and 'N/A' are not treated as degradation.
    """
    values = _fan_status_values(value)
    for item in values:
        ratio = fan_ratio(item)
        if ratio and ratio[0] < ratio[1]:
            return True
    return False


def degraded_fan_detail(fan_status):
    """
    Builds the 'Fan degraded: X/Y operational' detail for each degraded member.
    Returns '' when no member is degraded.
    """
    parts = []
    if isinstance(fan_status, dict):
        items = fan_status.items()
    else:
        items = [(None, v) for v in _fan_status_values(fan_status)]
    for member, item in items:
        ratio = fan_ratio(item)
        if ratio and ratio[0] < ratio[1]:
            where = f"member {member}: " if member is not None else ""
            parts.append("Fan degraded: {}{}/{} operational".format(where, ratio[0], ratio[1]))
    return "; ".join(parts)

def is_invalid_command_output(output):
    """
    True when the CLI rejected the command syntax rather than returning data.
    ArubaOS-CX answers `show system information` with an 'Invalid input' error.
    """
    if not output:
        return False
    lowered = output.lower()
    return any(marker in lowered for marker in INVALID_COMMAND_MARKERS)


def _output_has_data(collector, output):
    """
    True when a response actually carries data the collector can use.

    A platform may accept a command's syntax yet return nothing useful for it,
    so the fallback also triggers on an empty-but-valid response.
    """
    if not output or not output.strip():
        return False

    if collector == "sn":
        return bool(_extract_field(output, SERIAL_PATTERNS))
    if collector == "fan":
        return bool(_parse_cx_fan_table(output) or _extract_fan_ratio(output))
    if collector == "huawei-count":
        return "huawei" in output.lower() or "---" in output or "+" in output
    return True


def _run_with_fallback(net_connect, collector, cmd, read_timeout, logger=None):
    """
    Runs a collector command, retrying with the multi-OS equivalent when needed.

    ProCurve and ArubaOS-CX use different syntax for the same data, so each
    collector has a documented fallback:
      sn           -> 'show system'
      fan          -> 'show environment fan'
      huawei-count -> 'show lldp neighbor-info'
      poe          -> no fallback; a rejected command means 'Non-PoE'
    """
    def log(message, level="info"):
        if logger is not None:
            getattr(logger, level)(message)

    out = net_connect.send_command(cmd, read_timeout=read_timeout)

    if collector == "poe":
        if is_non_poe_output(out):
            log(f"[{collector}] '{cmd}' unsupported on this platform, recording {NON_POE}")
            return NON_POE
        return out

    fallback = COLLECTOR_FALLBACK_COMMANDS.get(collector)
    if not fallback:
        return out

    needs_fallback = (
        is_invalid_command_output(out)
        or not _output_has_data(collector, out)
    )
    if needs_fallback:
        log(f"[{collector}] '{cmd}' returned no usable data, retrying with '{fallback}'")
        out = net_connect.send_command(fallback, read_timeout=read_timeout)

    return out


def collect_device_data(record, device_types, read_timeout, mode, logger=None):
    """
    Establishes connection to the switch trying device_types in order.
    Executes relevant command(s) for the mode.
    Handles fallbacks and retries on timeout/transient errors.
    Fails fast on NetmikoAuthenticationException.
    Returns:
      {
        "status": "success" | "partial" | "failed",
        "outputs": { collector_name: raw_output },
        "errors": { collector_name: { "error_type": ..., "error_message": ..., "attempts": ... } },
        "attempts": { collector_name: current_attempt_count }
      }
    """
    collectors_to_run = get_collectors_to_run(mode)

    outputs = {}
    errors = {}
    attempts_dict = {c: 0 for c in collectors_to_run}

    success_collectors = set()

    if logger is None:
        # Create a no-op logger for backward compatibility
        logger = logging.getLogger("switch_collector.null")
        logger.addHandler(logging.NullHandler())

    logger.info(f"Starting connection attempts for {record['label']} ({record['ip']})")
    logger.info(f"Device types to try: {', '.join(device_types)} | Read timeout: {read_timeout}s | Mode: {mode}")

    for device_type in device_types:
        attempts = 0
        max_attempts = 3
        is_auth_failure = False

        logger.info(f"Trying device type: {device_type}")

        while attempts < max_attempts:
            attempts += 1
            attempted_this_session = set()

            logger.info(f"Connection attempt {attempts}/{max_attempts} with {device_type}")

            try:
                device = {
                    "device_type": device_type,
                    "host": record["ip"],
                    "username": record["username"],
                    "password": record["password"],
                    "conn_timeout": 10,
                    "auth_timeout": 10,
                    "banner_timeout": 10,
                }

                # Establish connection
                logger.info(f"Connecting to {record['ip']}...")
                with ConnectHandler(**device) as net_connect:
                    logger.info(f"Connected successfully to {record['ip']} via {device_type}")

                    # Run each required command
                    for c in collectors_to_run:
                        if c in success_collectors:
                            continue

                        cmd = COLLECTOR_COMMANDS.get(c, "")

                        attempts_dict[c] += 1
                        attempted_this_session.add(c)

                        logger.info(f"[{c}] Sending: '{cmd}' (attempt {attempts_dict[c]})")
                        try:
                            # Run command, retrying with the multi-OS
                            # equivalent when the platform rejects it.
                            out = _run_with_fallback(net_connect, c, cmd, read_timeout, logger)
                            outputs[c] = out
                            errors.pop(c, None)
                            success_collectors.add(c)
                            logger.info(f"[{c}] Success - received {len(out)} chars")
                        except NetmikoAuthenticationException:
                            raise
                        except (NetmikoTimeoutException, ReadTimeout, Exception) as e:
                            err_type = type(e).__name__
                            err_msg = str(e)
                            if record["password"] and (record["password"] in err_msg):
                                err_msg = err_msg.replace(record["password"], "********")
                            errors[c] = {
                                "error_type": err_type,
                                "error_message": err_msg,
                                "attempts": attempts_dict[c]
                            }
                            logger.warning(f"[{c}] Failed: {err_type}: {err_msg}")
                            continue

                # If all requested collectors ran successfully, return success!
                if len(success_collectors) == len(collectors_to_run):
                    logger.info(f"All collectors succeeded for {record['label']}")
                    return {
                        "status": "success",
                        "outputs": outputs,
                        "errors": errors,
                        "attempts": attempts_dict
                    }

            except NetmikoAuthenticationException as e:
                err_msg = str(e)
                if record["password"] and (record["password"] in err_msg):
                    err_msg = err_msg.replace(record["password"], "********")

                logger.error(f"Authentication failed: {err_msg}")

                # Record authentication failure for all uncompleted collectors.
                # If auth failed before any command, count this connection attempt.
                for c in collectors_to_run:
                    if c not in success_collectors:
                        if c not in attempted_this_session:
                            attempts_dict[c] += 1
                        errors[c] = {
                            "error_type": "NetmikoAuthenticationException",
                            "error_message": err_msg,
                            "attempts": attempts_dict[c]
                        }
                is_auth_failure = True
                break  # Fail-fast, do not retry auth errors

            except (NetmikoTimeoutException, ReadTimeout, Exception) as e:
                err_type = type(e).__name__
                err_msg = str(e)
                if record["password"] and (record["password"] in err_msg):
                    err_msg = err_msg.replace(record["password"], "********")

                logger.warning(f"Connection error: {err_type}: {err_msg}")

                # Connection-level failure before any command applies to all pending collectors.
                if not attempted_this_session:
                    for c in collectors_to_run:
                        if c not in success_collectors:
                            attempts_dict[c] += 1
                            errors[c] = {
                                "error_type": err_type,
                                "error_message": err_msg,
                                "attempts": attempts_dict[c]
                            }
                else:
                    for c in collectors_to_run:
                        if c not in success_collectors and c not in attempted_this_session:
                            errors[c] = {
                                "error_type": err_type,
                                "error_message": err_msg,
                                "attempts": attempts_dict[c]
                            }

        if is_auth_failure:
            logger.info("Stopping driver fallback due to authentication failure")
            break  # Exit device_type fallback on authentication failures

        if len(success_collectors) == len(collectors_to_run):
            logger.info("All collectors completed, stopping driver fallback")
            break  # Stop trying other device types if we completed all collectors

        logger.info("Switching to next device type")

    # Compute overall SSH status
    if len(success_collectors) == len(collectors_to_run):
        status = "success"
    elif success_collectors:
        status = "partial"
    else:
        status = "failed"

    logger.info(f"Final status for {record['label']}: {status.upper()} (success: {len(success_collectors)}/{len(collectors_to_run)})")
    if errors:
        for c, err in errors.items():
            logger.info(f"  [{c}] {err['error_type']}: {err['error_message']} (attempts: {err['attempts']})")

    return {
        "status": status,
        "outputs": outputs,
        "errors": errors,
        "attempts": attempts_dict
    }

def process_parsed_data(outputs, label, collectors_to_run, attempts_dict, errors_dict):
    """
    Parses outputs. If a parser raises ValueError, records ParseError in errors_dict
    and removes from outputs or marks as failed.
    """
    parsed_data = {}
    for c in collectors_to_run:
        if c in outputs:
            try:
                if c == "sn":
                    parsed_data["sn"] = parse_system_information(outputs["sn"], label=label)
                elif c == "fan":
                    parsed_data["fan"] = parse_fan_output(outputs["fan"])
                elif c == "poe":
                    parsed_data["poe"] = parse_power_output(outputs["poe"])
                elif c == "huawei-count":
                    parsed_data["huawei-count"] = count_huawei_devices(outputs["huawei-count"])
            except ValueError as e:
                # ParseError triggers fail-fast on parsing. No SSH retry!
                errors_dict[c] = {
                    "error_type": "ParseError",
                    "error_message": str(e),
                    "attempts": attempts_dict[c]
                }
                outputs.pop(c, None)
    return parsed_data

def setup_logging(log_dir: str) -> tuple[logging.Logger, str]:
    """
    Set up file logging with timestamped log file.
    Returns (logger, log_file_path).
    """
    Path(log_dir).mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    log_file = os.path.join(log_dir, f"collector_{timestamp}.log")

    logger = logging.getLogger("switch_collector")
    logger.setLevel(logging.DEBUG)

    # File handler - detailed debug logs
    file_handler = logging.FileHandler(log_file, encoding='utf-8')
    file_handler.setLevel(logging.DEBUG)
    file_format = logging.Formatter(
        '%(asctime)s | %(levelname)-8s | %(message)s',
        datefmt='%Y-%m-%d %H:%M:%S'
    )
    file_handler.setFormatter(file_format)
    logger.addHandler(file_handler)

    # Prevent propagation to root logger
    logger.propagate = False

    return logger, log_file


def process_switch(record, device_types, read_timeout, mode, logger):
    """
    Process a single switch - wrapper for ThreadPoolExecutor.
    Returns a tuple of (record, raw_res, parsed_res, final_errors).
    """
    raw_res = collect_device_data(record, device_types, read_timeout, mode, logger)

    # Parse raw outputs
    parsed_res = process_parsed_data(
        outputs=raw_res["outputs"],
        label=record["label"],
        collectors_to_run=get_collectors_to_run(mode),
        attempts_dict=raw_res["attempts"],
        errors_dict=raw_res["errors"]
    )

    final_errors = raw_res["errors"]
    return record, raw_res, parsed_res, final_errors


def main():
    args = parse_args()
    mode, workers = resolve_runtime_options(args)
    if mode is None:
        print("Exiting. Goodbye!")
        raise SystemExit(0)

    device_types = [dt.strip() for p in args.device_types.split(",") for dt in p.split("|") if dt.strip()]
    read_timeout = args.read_timeout

    # BASE_DIR is the folder holding the executable (frozen) or the script.
    inventory_path = args.inventory or str(BASE_DIR / "switch.txt")
    output_dir = args.output_dir or str(BASE_DIR / "outputs")
    # A relative --log-dir is resolved against BASE_DIR so a double-clicked
    # executable still writes next to itself, not into an arbitrary CWD.
    log_dir = args.log_dir
    if not os.path.isabs(log_dir):
        log_dir = str(BASE_DIR / log_dir)

    os.makedirs(output_dir, exist_ok=True)

    # Set up logging
    logger, log_file = setup_logging(log_dir)

    print(f"Loading inventory from: {inventory_path}")
    try:
        records, parse_errors = parse_inventory(inventory_path)
    except FileNotFoundError as e:
        print(f"Error: {e}")
        return

    total_switches = len(records)
    collectors_to_run = get_collectors_to_run(mode)

    # Startup summary
    print(f"Switches: {total_switches} | Mode: {mode} | Workers: {workers} | Log: {log_file}")
    logger.info(f"Starting collection: switches={total_switches}, mode={mode}, workers={workers}, device_types={device_types}")

    # Unified Excel rows. Aggregation happens in the main thread, so worker
    # threads never touch shared state.
    excel_rows = []

    # Inventory lines that never became a device get one failed row each.
    for p_err in parse_errors:
        excel_rows.append(build_failed_row(
            p_err["label"], p_err["ip"],
            error_type=p_err["error_type"],
            error_msg=p_err["error_message"],
        ))

    start_time = time.time()
    completed = 0

    if workers == 1:
        # Sequential execution
        for record in records:
            completed += 1
            logger.info(f"[{completed}/{total_switches}] Processing {record['label']} ({record['ip']})")
            record, raw_res, parsed_res, final_errors = process_switch(record, device_types, read_timeout, mode, logger)
            excel_rows.extend(build_member_rows(
                record, parsed_res, final_errors, mode, collectors_to_run))
            _print_progress(completed, total_switches, record, raw_res["status"], final_errors, mode)
    else:
        # Parallel execution
        with ThreadPoolExecutor(max_workers=workers) as executor:
            future_to_record = {
                executor.submit(process_switch, record, device_types, read_timeout, mode, logger): record
                for record in records
            }

            for future in as_completed(future_to_record):
                completed += 1
                record = future_to_record[future]
                try:
                    record, raw_res, parsed_res, final_errors = future.result()
                    excel_rows.extend(build_member_rows(
                        record, parsed_res, final_errors, mode, collectors_to_run))
                    _print_progress(completed, total_switches, record, raw_res["status"], final_errors, mode)
                except Exception as e:
                    logger.error(f"Execution error for {record['label']}: {e}")
                    excel_rows.append(build_failed_row(
                        record["label"], record["ip"],
                        error_type=type(e).__name__, error_msg=str(e),
                    ))
                    _print_progress(completed, total_switches, record, "failed", {"error": str(e)}, mode)

    elapsed = time.time() - start_time

    # Single consolidated Excel report, named for the audit type.
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    excel_path = build_report_path(output_dir, mode, timestamp)
    write_excel_file(excel_rows, excel_path, mode=mode)

    # Summary counts are per switch, so a stack is counted once, not per member.
    switch_status = {}
    for row in excel_rows:
        key = (row["Label"], row["IP Address"])
        # A switch is only as healthy as its best member row.
        current = switch_status.get(key)
        if current is None or STATUS_RANK.get(row["Status"], 3) < STATUS_RANK.get(current, 3):
            switch_status[key] = row["Status"]

    success_count = sum(1 for s in switch_status.values() if s == "success")
    warning_count = sum(1 for s in switch_status.values() if s == "warning")
    partial_count = sum(1 for s in switch_status.values() if s == "partial")
    failed_count = sum(1 for s in switch_status.values() if s == "failed")

    print(f"Results written to: {excel_path}")
    print(f"\nCompleted in {elapsed:.1f}s | {success_count} success, {warning_count} warning, "
          f"{partial_count} partial, {failed_count} failed | Results in {output_dir}/")
    logger.info(f"Collection complete: {elapsed:.1f}s, success={success_count}, "
                f"warning={warning_count}, partial={partial_count}, failed={failed_count}")


# Column sets per mode so a single-mode report never shows empty columns.
_COMMON_TAIL = ["Status", "Error Detail", "Solution Hint"]
_IDENTITY = ["No", "Label", "IP Address", "Hostname", "Member"]

EXCEL_COLUMNS_BY_MODE = {
    "all": _IDENTITY + [
        "Serial Number", "ROM Version", "FAN",
        "PoE Remaining (W)", "Huawei AP Count",
    ] + _COMMON_TAIL,
    "fan": _IDENTITY + ["FAN"] + _COMMON_TAIL,
    "sn": _IDENTITY + ["Serial Number", "ROM Version"] + _COMMON_TAIL,
    "poe": _IDENTITY + ["PoE Remaining (W)"] + _COMMON_TAIL,
    "huawei-count": _IDENTITY + ["Huawei AP Count"] + _COMMON_TAIL,
}

# Default (full 'all') column set, kept for callers that do not pass a mode.
EXCEL_COLUMNS = EXCEL_COLUMNS_BY_MODE["all"]

# Report filename stem per mode.
REPORT_STEMS = {
    "all": "switch_master",
    "fan": "switch_fan_audit",
    "sn": "switch_asset_inventory",
    "poe": "switch_poe_capacity",
    "huawei-count": "switch_huawei_ap",
}


def get_columns_for_mode(mode):
    """
    Returns the Excel column list for a mode, falling back to the full set.
    """
    return EXCEL_COLUMNS_BY_MODE.get(mode, EXCEL_COLUMNS_BY_MODE["all"])


def build_report_path(output_dir, mode, timestamp):
    """
    Builds the mode-specific report path, e.g.
    outputs/switch_fan_audit_20260928_093000.xlsx
    """
    stem = REPORT_STEMS.get(mode, "switch_master")
    return os.path.join(output_dir, "{}_{}.xlsx".format(stem, timestamp))

# Status cell styling: fill / font colour.
STATUS_STYLES = {
    "success": ("C6EFCE", "006100"),
    "warning": ("FFEB9C", "9C6500"),
    "partial": ("FFEB9C", "9C6500"),
    "failed": ("FFC7CE", "9C0006"),
}

# Severity ordering when a switch has several member rows.
STATUS_RANK = {"success": 0, "warning": 1, "partial": 2, "failed": 3}


def _poe_to_watts(poe_value):
    """'512 W' -> '512'. The column header already carries the unit."""
    if not poe_value:
        return ""
    match = re.search(r'(\d+(?:\.\d+)?)', str(poe_value))
    return match.group(1) if match else str(poe_value)


def _ordered_member_ids(sn_members, fan, huawei):
    """
    Returns the ordered VSF member ids to render as rows.

    Real member ids come from the fan / Huawei dicts when available; otherwise
    members are numbered positionally from the parsed SN list.
    """
    ids = set()
    for source in (fan, huawei):
        if isinstance(source, dict):
            for key in source:
                if str(key).isdigit():
                    ids.add(int(key))
    if ids:
        return [str(i) for i in sorted(ids)]
    if sn_members:
        return [str(i + 1) for i in range(len(sn_members))]
    return ["1"]


def _device_status(parsed_res, final_errors, collectors_to_run):
    """Returns (status, error_detail) for a device across its collectors."""
    failed = [c for c in collectors_to_run if c not in parsed_res]
    if not failed:
        return "success", ""
    status = "partial" if any(c in parsed_res for c in collectors_to_run) else "failed"
    detail = "; ".join(
        "{}={}: {}".format(
            c,
            final_errors.get(c, {}).get("error_type", "Unknown"),
            # Netmiko errors are often multi-line; keep the cell on one line so
            # the sheet stays readable. The full text is preserved in the log.
            " ".join(str(final_errors.get(c, {}).get("error_message", "")).split()),
        )
        for c in failed
    )
    return status, detail


def _solution_hint_for(final_errors, failed_collectors, fan_value, status):
    """
    First matching remediation hint for a single member row.

    Collector failures are the strongest signal, then this member's own fan
    health, then a generic partial-failure hint.
    """
    for c in failed_collectors:
        err = final_errors.get(c, {})
        hint = get_solution_hint(err.get("error_type", ""), err.get("error_message", ""), None)
        if hint:
            return hint

    if fan_is_degraded(fan_value):
        return HINT_FAN

    if status in ("partial", "warning"):
        return HINT_PARTIAL
    return ""


def build_failed_row(label, ip, error_type, error_msg):
    """A single failed row (member 1) for a device that produced no data."""
    return {
        "No": 0,
        "Label": label,
        "IP Address": ip,
        "Hostname": "",
        "Member": "1",
        "Serial Number": "",
        "ROM Version": "",
        "FAN": "",
        "PoE Remaining (W)": "",
        "Huawei AP Count": "",
        "Status": "failed",
        "Error Detail": "{}: {}".format(error_type, " ".join(str(error_msg).split())),
        "Solution Hint": get_solution_hint(error_type, error_msg, None),
    }


def build_member_rows(record, parsed_res, final_errors, mode, collectors_to_run):
    """
    Unrolls one switch into one Excel row per physical member.

    Standalone switches produce a single row (Member 1). VSF stacks produce one
    row per member with that member's SN, ROM, fan state and Huawei AP count.
    PoE remaining power is shared across the whole stack.
    """
    sn_members = parsed_res.get("sn", [])
    fan = parsed_res.get("fan", {})
    huawei = parsed_res.get("huawei-count", {})
    poe = _poe_to_watts(parsed_res.get("poe", ""))

    status, error_detail = _device_status(parsed_res, final_errors, collectors_to_run)
    failed_collectors = [c for c in collectors_to_run if c not in parsed_res]

    if "sn" in collectors_to_run and not sn_members:
        # SN could not be read, so there is no member layout to unroll.
        return [build_failed_row(
            record["label"], record["ip"],
            error_type=final_errors.get("sn", {}).get("error_type", "Unknown"),
            error_msg=final_errors.get("sn", {}).get("error_message", ""),
        )]

    member_ids = _ordered_member_ids(sn_members, fan, huawei)
    hostname = sn_members[0]["hostname"] if sn_members else ""

    rows = []
    for index, member_id in enumerate(member_ids):
        member = sn_members[index] if index < len(sn_members) else {}
        fan_value = fan.get(member_id, "") if isinstance(fan, dict) else ""

        # A degraded fan downgrades an otherwise-successful row to 'warning'.
        member_degraded = fan_is_degraded(fan_value)
        row_status = "warning" if (member_degraded and status == "success") else status
        row_detail = error_detail
        if member_degraded:
            fan_detail = "Fan degraded: {}/{} operational".format(*fan_ratio(fan_value))
            row_detail = "{}; {}".format(fan_detail, error_detail) if error_detail else fan_detail

        # Prefer a collector failure hint; fall back to this member's fan state.
        hint = _solution_hint_for(final_errors, failed_collectors, fan_value, row_status)

        rows.append({
            "No": 0,
            "Label": record["label"],
            "IP Address": record["ip"],
            "Hostname": member.get("hostname", hostname),
            "Member": member_id,
            "Serial Number": member.get("serial_number", ""),
            "ROM Version": member.get("rom_version", ""),
            "FAN": fan_value,
            "PoE Remaining (W)": poe,
            "Huawei AP Count": huawei.get(member_id, 0) if isinstance(huawei, dict) else 0,
            "Status": row_status,
            "Error Detail": row_detail,
            "Solution Hint": hint,
        })
    return rows


def write_excel_file(rows, filepath, mode=None, columns=None):
    """
    Writes the consolidated report as a single-sheet .xlsx workbook.

    One worksheet named 'Switch Data', a frozen dark-navy header, colour-coded
    status cells, and auto-fitted column widths. Only the columns relevant to
    `mode` are written, so a single-mode report has no empty columns.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    if columns is None:
        columns = get_columns_for_mode(mode)
    # Never emit a sheet without a status column to colour.
    if "Status" not in columns:
        columns = list(columns) + ["Status"]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Switch Data"

    sheet.append(list(columns))

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    sheet.freeze_panes = "A2"

    status_index = columns.index("Status") + 1

    for position, row in enumerate(rows, start=1):
        values = dict(row)
        values["No"] = position
        sheet.append([values.get(column, "") for column in columns])

        fill_colour, font_colour = STATUS_STYLES.get(row.get("Status", ""), (None, None))
        if fill_colour:
            status_cell = sheet.cell(row=sheet.max_row, column=status_index)
            status_cell.fill = PatternFill("solid", fgColor=fill_colour)
            status_cell.font = Font(color=font_colour)

    # Auto-fit column widths with a small padding.
    for column_index, column_cells in enumerate(sheet.columns, start=1):
        width = 0
        for cell in column_cells:
            if cell.value is not None:
                width = max(width, len(str(cell.value)))
        sheet.column_dimensions[get_column_letter(column_index)].width = width + 4

    workbook.save(filepath)
    return filepath


def _print_progress(completed, total, record, status, errors, mode):
    """Print compact progress line for console."""
    if status == "success":
        print(f"[{completed}/{total}] {record['label']} ({record['ip']}) -> [OK]")
    elif status == "warning":
        print(f"[{completed}/{total}] {record['label']} ({record['ip']}) -> [WARNING]: Fan degraded")
    elif status == "partial":
        print(f"[{completed}/{total}] {record['label']} ({record['ip']}) -> [PARTIAL]")
    else:
        # First collector error becomes the one-line console summary; the full
        # multi-line text stays in the log file.
        if isinstance(errors, dict) and errors:
            first_err = next(iter(errors.values()))
            if isinstance(first_err, dict):
                err_msg = first_err.get("error_message", str(first_err))
            else:
                err_msg = str(first_err)
        else:
            err_msg = str(errors) if errors else "Unknown error"
        # Collapse to a single line and keep it short.
        err_msg = " ".join(str(err_msg).split())
        err_msg = err_msg[:80] + "..." if len(err_msg) > 80 else err_msg
        print(f"[{completed}/{total}] {record['label']} ({record['ip']}) -> [FAIL]: {err_msg}")

if __name__ == "__main__":
    main()
